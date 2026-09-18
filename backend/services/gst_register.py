"""HQ GST register — every tenant payment as an invoice line, month-wise, exportable to Excel for the CA."""
import io
from collections import defaultdict

from database import _raw_db
from services.hq_tax import get_profile, tax_breakdown
from services.tax_invoice import _invoice_number


async def register_rows(month: str | None = None) -> list[dict]:
    """month = 'YYYY-MM' (IST calendar month) or None for everything."""
    hq = await get_profile()
    tenants = {t["id"]: t async for t in _raw_db.tenants.find({}, {"_id": 0, "id": 1, "name": 1, "slug": 1, "gstin": 1, "location": 1, "state_code": 1})}
    hq_state = (hq.get("gstin") or hq.get("state_code") or "29")[:2]
    rows = []
    sources = ((_raw_db.subscription_payments, {"kind": {"$ne": "razorpay_pending"}}, "Subscription"),
               (_raw_db.sms_pack_payments, {"status": "captured"}, "Message credits"))
    for coll, q, kind in sources:
        async for p in coll.find(q, {"_id": 0}):
            paid_at = (p.get("paid_at") or p.get("captured_at") or p.get("created_at") or "")[:10]
            if not paid_at or (month and not paid_at.startswith(month)):
                continue
            t = tenants.get(p.get("tenant_id")) or {}
            tax = p.get("tax") or tax_breakdown(float(p.get("amount") or 0) / (1 + float(hq.get("gst_rate_pct") or 0) / 100), hq)
            same_state = (t.get("gstin") or "")[:2] == hq_state if t.get("gstin") else True
            gst = float(tax.get("gst") or 0)
            cgst = round(gst / 2, 2) if same_state else 0.0
            rows.append({
                "invoice_no": p.get("invoice_no") or await _invoice_number(p, coll), "date": paid_at, "kind": kind,
                "description": (f"Plan {p.get('plan') or ''}".strip() if kind == "Subscription" else f"{p.get('points')} {p.get('channel') or 'sms'} credits"),
                "tenant": t.get("name") or p.get("tenant_id") or "", "tenant_gstin": t.get("gstin") or "", "place_of_supply": "Karnataka (29)" if same_state else "Other state",
                "sac": "998314", "taxable": float(tax.get("base") or 0), "rate": float(tax.get("gst_rate_pct") or 0),
                "cgst": cgst, "sgst": round(gst - cgst, 2) if same_state else 0.0, "igst": 0.0 if same_state else round(gst, 2),
                "total": float(tax.get("total") or p.get("amount") or 0), "payment_ref": p.get("txn_ref") or p.get("razorpay_payment_id") or "", "method": p.get("method") or "razorpay",
            })
    rows.sort(key=lambda r: (r["date"], r["invoice_no"]))
    return rows


def month_summary(rows: list[dict]) -> list[dict]:
    agg = defaultdict(lambda: {"invoices": 0, "taxable": 0.0, "cgst": 0.0, "sgst": 0.0, "igst": 0.0, "total": 0.0})
    for r in rows:
        m = agg[r["date"][:7]]
        m["invoices"] += 1
        for k in ("taxable", "cgst", "sgst", "igst", "total"):
            m[k] = round(m[k] + r[k], 2)
    return [{"month": k, **v} for k, v in sorted(agg.items(), reverse=True)]


def to_xlsx(rows: list[dict], month: str | None, hq: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "GST Register"
    title = f"Miracurl Studio — GST Sales Register {month or 'all months'} · GSTIN {hq.get('gstin') or 'pending'} · MSME {hq.get('msme') or '-'}"
    ws.append([title]); ws["A1"].font = Font(bold=True, size=12); ws.append([])
    heads = ["Invoice No", "Date", "Type", "Description", "Customer", "Customer GSTIN", "Place of Supply", "SAC", "Taxable Value", "GST %", "CGST", "SGST", "IGST", "Invoice Total", "Payment Ref", "Method"]
    ws.append(heads)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="14100A"); c.alignment = Alignment(horizontal="center")
    keys = ["invoice_no", "date", "kind", "description", "tenant", "tenant_gstin", "place_of_supply", "sac", "taxable", "rate", "cgst", "sgst", "igst", "total", "payment_ref", "method"]
    for r in rows:
        ws.append([r[k] for k in keys])
    n = len(rows)
    if n:
        tot_row = ["TOTAL", "", "", "", "", "", "", "", f"=SUM(I4:I{3 + n})", "", f"=SUM(K4:K{3 + n})", f"=SUM(L4:L{3 + n})", f"=SUM(M4:M{3 + n})", f"=SUM(N4:N{3 + n})", "", ""]
        ws.append(tot_row)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for i, w in enumerate([18, 11, 14, 28, 26, 17, 16, 8, 14, 7, 11, 11, 11, 14, 22, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=9, max_col=14):
        for c in row:
            c.number_format = "#,##0.00"
    ws.freeze_panes = "A4"
    ws2 = wb.create_sheet("Monthly summary"); ws2.append(["Month", "Invoices", "Taxable Value", "CGST", "SGST", "IGST", "Total"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for m in month_summary(rows):
        ws2.append([m["month"], m["invoices"], m["taxable"], m["cgst"], m["sgst"], m["igst"], m["total"]])
    for i, w in enumerate([10, 10, 14, 12, 12, 12, 14], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
